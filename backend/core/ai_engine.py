"""
Groq AI Engine
- STT: Whisper qua Groq API (nhanh, miễn phí)
- LLM: Llama3 qua Groq API (nhanh, miễn phí)
- RAG: ChromaDB + sentence-transformers (local)
"""
import os, re, json, asyncio, tempfile, logging
from pathlib import Path
from typing import Optional, List, Dict

log = logging.getLogger("groq_engine")

GROQ_KEY  = os.getenv("GROQ_API_KEY", "")
LLM_MODEL = os.getenv("GROQ_LLM_MODEL", "llama-3.1-8b-instant")
STT_MODEL = os.getenv("GROQ_STT_MODEL", "whisper-large-v3-turbo")
COMPANY   = os.getenv("COMPANY_NAME", "Công ty")
COMPANY_D = os.getenv("COMPANY_DESC", "")
AI_NAME   = os.getenv("AI_NAME", "Aria")

SYSTEM_PROMPT = """Bạn là {name}, AI Lễ Tân của {company}. {desc}

QUY TẮC QUAN TRỌNG:
- Trả lời tiếng Việt, thân thiện, ngắn gọn (1-3 câu)
- Xưng "em", gọi khách "anh/chị"
- KHÔNG dùng emoji, icon, bullet points, markdown (không dùng *, **, #, -)
- KHÔNG viết danh sách, chỉ viết câu văn bình thường
- Sau 2-3 câu khách chưa nói tên thì hỏi tên tự nhiên
- Dùng tên khách nếu đã biết
- KHÔNG đề cập bạn là AI trừ khi được hỏi trực tiếp

TRẠNG THÁI: {state}
KHÁCH: {guest}
KIẾN THỨC CÔNG TY:
{context}

LỊCH SỬ HỘI THOẠI:
{history}"""

ROBOT_MAP = {
    "robot công nghiệp": "industrial", "robot hàn": "welding",
    "cobot": "cobot", "robot kho": "warehouse",
    "robot dịch vụ": "service", "cánh tay robot": "arm",
    "robot lắp ráp": "assembly", "robot lễ tân": "service",
}


# ── Chat Session ────────────────────────────────────────────────────
class ChatSession:
    def __init__(self, guest_id=None, guest_name=None):
        self.guest_id   = guest_id
        self.guest_name = guest_name
        self.history: List[Dict] = []
        self.turn       = 0
        self.name_asked = False
        self.state      = "greeting"

    def add(self, role: str, text: str):
        self.history.append({"role": role, "content": text})
        if role == "user":
            self.turn += 1
        if len(self.history) > 20:
            self.history = self.history[-20:]

    def history_str(self) -> str:
        lines = []
        for m in self.history[-8:]:
            who = "Khách" if m["role"] == "user" else AI_NAME
            lines.append(f"{who}: {m['content']}")
        return "\n".join(lines)

    def should_ask_name(self) -> bool:
        return (self.guest_name is None
                and self.turn >= 2
                and not self.name_asked
                and self.state not in ("booking", "contacting"))


# ── Knowledge Base ──────────────────────────────────────────────────
class KnowledgeBase:
    def __init__(self):
        self._col  = None
        self._emb  = None
        self._init()

    def _init(self):
        try:
            import chromadb
            path = os.getenv("CHROMA_PATH", "./data/chroma_db")
            client = chromadb.PersistentClient(path=path)
            self._col = client.get_or_create_collection(
                "aria_kb", metadata={"hnsw:space": "cosine"})
            log.info(f"ChromaDB ready — {self._col.count()} chunks")
        except Exception as e:
            log.error(f"ChromaDB init failed: {e}")

    def _embedder(self):
        if self._emb is None:
            try:
                from sentence_transformers import SentenceTransformer
                self._emb = SentenceTransformer(
                    "intfloat/multilingual-e5-base", device="cpu")
                log.info("Embedder loaded")
            except Exception as e:
                log.error(f"Embedder failed: {e}")
        return self._emb

    async def query(self, text: str, n: int = 4) -> str:
        if not self._col or self._col.count() == 0:
            return ""
        emb = self._embedder()
        if emb is None:
            return ""
        loop = asyncio.get_event_loop()
        vec  = await loop.run_in_executor(None, lambda: emb.encode(text).tolist())
        try:
            res  = self._col.query(
                query_embeddings=[vec],
                n_results=min(n, self._col.count()),
                include=["documents"])
            docs = res.get("documents", [[]])[0]
            return "\n---\n".join(docs) if docs else ""
        except Exception as e:
            log.error(f"KB query error: {e}")
            return ""

    async def ingest_file(self, path: str, doc_id: str) -> int:
        text = await self._read_file(path)
        if not text.strip():
            return 0
        emb = self._embedder()
        if emb is None or self._col is None:
            return 0
        chunks = self._chunk(text)
        loop   = asyncio.get_event_loop()
        for i, chunk in enumerate(chunks):
            vec = await loop.run_in_executor(
                None, lambda c=chunk: emb.encode(c).tolist())
            cid = f"{doc_id}_c{i}"
            try:
                self._col.add(ids=[cid], documents=[chunk], embeddings=[vec],
                              metadatas=[{"doc_id": doc_id, "i": i}])
            except Exception:
                pass  # chunk đã tồn tại
        return len(chunks)

    async def _read_file(self, path: str) -> str:
        ext = Path(path).suffix.lower()
        loop = asyncio.get_event_loop()
        try:
            if ext == ".pdf":
                return await loop.run_in_executor(None, self._read_pdf, path)
            elif ext in (".docx", ".doc"):
                return await loop.run_in_executor(None, self._read_docx, path)
            elif ext in (".xlsx", ".xls"):
                return await loop.run_in_executor(None, self._read_excel, path)
            elif ext in (".txt", ".md"):
                return Path(path).read_text(encoding="utf-8", errors="ignore")
        except Exception as e:
            log.error(f"Read file error {path}: {e}")
        return ""

    def _read_pdf(self, path: str) -> str:
        import fitz
        doc  = fitz.open(path)
        return "\n".join(p.get_text() for p in doc)

    def _read_docx(self, path: str) -> str:
        from docx import Document
        return "\n".join(p.text for p in Document(path).paragraphs if p.text.strip())

    def _read_excel(self, path: str) -> str:
        import openpyxl
        wb   = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                r = " | ".join(str(c) for c in row if c is not None)
                if r.strip():
                    rows.append(r)
        return "\n".join(rows)

    def _chunk(self, text: str, size: int = 500, overlap: int = 50) -> List[str]:
        if len(text) <= size:
            return [text]
        chunks, start = [], 0
        while start < len(text):
            end   = start + size
            chunk = text[start:end]
            for sep in ("\n\n", "\n", ". ", ", "):
                last = chunk.rfind(sep)
                if last > size // 2:
                    chunk = chunk[:last + len(sep)]
                    break
            chunks.append(chunk.strip())
            start += len(chunk) - overlap
        return [c for c in chunks if c.strip()]


# ── Groq LLM ────────────────────────────────────────────────────────
class GroqLLM:
    def __init__(self, kb: KnowledgeBase):
        self.kb = kb

    async def respond(self, session: ChatSession, user_text: str) -> dict:
        session.add("user", user_text)

        # Hỏi tên nếu cần
        if session.should_ask_name():
            session.name_asked = True
            session.state = "asking_name"

        # Trích xuất tên từ câu trả lời
        if session.state == "asking_name" and not session.guest_name:
            name = self._extract_name(user_text)
            if name:
                session.guest_name = name
                session.state = "chatting"

        # Phát hiện intent
        intent = self._intent(user_text)
        robot  = self._robot_model(user_text)
        if intent in ("booking", "contacting"):
            session.state = intent

        # RAG
        context = await self.kb.query(user_text)

        # Build system prompt
        state_desc = {
            "greeting":    "Vừa gặp khách, chào thân thiện",
            "chatting":    "Đang trò chuyện",
            "asking_name": "Cần hỏi tên khách tự nhiên",
            "booking":     "Đang hỗ trợ đặt lịch",
            "contacting":  "Đang liên hệ nhân viên",
        }.get(session.state, "Trò chuyện")

        sys_prompt = SYSTEM_PROMPT.format(
            name    = AI_NAME,
            company = COMPANY,
            desc    = COMPANY_D,
            state   = state_desc,
            guest   = f"Tên: {session.guest_name}" if session.guest_name else "Chưa biết tên",
            context = context or "Không có thông tin đặc biệt.",
            history = session.history_str()
        )

        reply = await self._call_groq(sys_prompt, user_text)
        session.add("assistant", reply)

        return {
            "text":       reply,
            "intent":     intent,
            "robot_model": robot,
            "guest_name": session.guest_name,
        }

    async def _call_groq(self, system: str, user: str) -> str:
        if not GROQ_KEY:
            return "Chưa cài GROQ_API_KEY. Vui lòng kiểm tra file .env."
        import httpx
        try:
            async with httpx.AsyncClient(timeout=20.0) as c:
                r = await c.post(
                    "https://api.groq.com/openai/v1/chat/completions",
                    headers={"Authorization": f"Bearer {GROQ_KEY}",
                             "Content-Type": "application/json"},
                    json={
                        "model": LLM_MODEL,
                        "messages": [
                            {"role": "system", "content": system},
                            {"role": "user",   "content": user}
                        ],
                        "max_tokens": 200,
                        "temperature": 0.7,
                    }
                )
                r.raise_for_status()
                return r.json()["choices"][0]["message"]["content"].strip()
        except Exception as e:
            log.error(f"Groq LLM error: {e}")
            return "Xin lỗi, em đang gặp sự cố nhỏ. Anh/chị vui lòng hỏi lại nhé."

    def _intent(self, text: str) -> str:
        t = text.lower()
        if any(k in t for k in ["đặt lịch", "hẹn gặp", "lịch hẹn", "hẹn"]):
            return "booking"
        if any(k in t for k in ["gặp anh", "gặp chị", "liên hệ", "tìm", "phòng ban"]):
            return "contacting"
        if any(k in t for k in ["robot", "sản phẩm", "xem", "giới thiệu", "mẫu"]):
            return "show_robot"
        if any(k in t for k in ["tạm biệt", "bye", "cảm ơn", "về rồi"]):
            return "farewell"
        return "chat"

    def _robot_model(self, text: str) -> Optional[str]:
        t = text.lower()
        for kw, mid in ROBOT_MAP.items():
            if kw in t:
                return mid
        if "robot" in t and any(k in t for k in ["xem", "giới thiệu", "mẫu", "loại"]):
            return "catalog"
        return None

    def _extract_name(self, text: str) -> Optional[str]:
        patterns = [
            r"(?:tên|tôi|em|mình)\s+(?:là|tên)\s+"
            r"([A-ZÀÁÂÃÈÉÊÌÍÒÓÔÕÙÚĂĐ][a-zàáâãèéêìíòóôõùúăđ]+"
            r"(?:\s+[A-ZÀÁÂÃÈÉÊÌÍÒÓÔÕÙÚĂĐ][a-zàáâãèéêìíòóôõùúăđ]+)*)",
        ]
        for p in patterns:
            m = re.search(p, text, re.IGNORECASE | re.UNICODE)
            if m:
                return m.group(1).strip().title()
        return None


# ── Groq STT ────────────────────────────────────────────────────────
class GroqSTT:
    async def transcribe(self, audio_bytes: bytes) -> str:
        if not GROQ_KEY:
            return ""
        import httpx
        try:
            with tempfile.NamedTemporaryFile(suffix=".webm", delete=False) as f:
                f.write(audio_bytes)
                tmp = f.name

            async with httpx.AsyncClient(timeout=30.0) as c:
                with open(tmp, "rb") as audio_file:
                    r = await c.post(
                        "https://api.groq.com/openai/v1/audio/transcriptions",
                        headers={"Authorization": f"Bearer {GROQ_KEY}"},
                        data={"model": STT_MODEL, "language": "vi",
                              "response_format": "text"},
                        files={"file": ("audio.webm", audio_file, "audio/webm")}
                    )
                    r.raise_for_status()
                    text = r.text.strip()

            import os; os.unlink(tmp)
            return text
        except Exception as e:
            log.error(f"Groq STT error: {e}")
            return ""


# ── TTS (gTTS) ──────────────────────────────────────────────────────
def _clean_for_tts(text: str) -> str:
    """Làm sạch text trước khi đưa vào TTS — loại bỏ emoji, markdown, ký tự lạ"""
    import unicodedata
    # Xóa emoji và ký tự unicode đặc biệt
    cleaned = ""
    for ch in text:
        cat = unicodedata.category(ch)
        # Giữ lại: chữ cái (L), số (N), dấu câu (P), khoảng trắng (Z), dấu thanh tiếng Việt (M)
        if cat.startswith(('L', 'N', 'P', 'Z', 'M')):
            cleaned += ch
    # Xóa markdown: **, *, _, ~, `, #
    cleaned = re.sub(r'[*_~`#>]', '', cleaned)
    # Xóa URL
    cleaned = re.sub(r'https?://\S+', '', cleaned)
    # Xóa nhiều khoảng trắng
    cleaned = re.sub(r'\s+', ' ', cleaned).strip()
    return cleaned


class TTSEngine:
    def __init__(self):
        self._cache_dir = Path("data/tts_cache")
        self._cache_dir.mkdir(parents=True, exist_ok=True)
        self._speed  = float(os.getenv("TTS_SPEED", "1.0"))
        self._lang   = os.getenv("TTS_LANG", "vi")

    async def speak(self, text: str) -> Optional[bytes]:
        if not text.strip():
            return None
        # Làm sạch text trước
        text = _clean_for_tts(text)
        if not text.strip():
            return None
        import hashlib
        key   = hashlib.md5(f"{text}{self._lang}".encode()).hexdigest()
        cache = self._cache_dir / f"{key}.mp3"
        if cache.exists():
            return cache.read_bytes()
        loop = asyncio.get_event_loop()
        data = await loop.run_in_executor(None, self._synth, text)
        if data:
            cache.write_bytes(data)
        return data

    def _synth(self, text: str) -> Optional[bytes]:
        try:
            from gtts import gTTS
            import io
            slow = self._speed < 0.85
            buf  = io.BytesIO()
            gTTS(text=text, lang=self._lang, slow=slow).write_to_fp(buf)
            return buf.getvalue()
        except Exception as e:
            log.error(f"gTTS error: {e}")
            return None

    def update(self, speed: float = None, lang: str = None):
        if speed is not None:
            self._speed = speed
            for f in self._cache_dir.glob("*.mp3"):
                f.unlink()
        if lang is not None:
            self._lang = lang
